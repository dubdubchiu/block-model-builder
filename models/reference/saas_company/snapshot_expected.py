"""Snapshot the reference workbook's cached values into expected.json.

The Expected sheets link live to the calculation rows, so an edit in Excel moves
them. Tests read expected.json, never the workbook. Re-run this only when the
reference model is deliberately changed, and commit both files together.

    uv run --no-project --python 3.12 --with openpyxl python snapshot_expected.py
"""
import datetime as dt
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).parent
WORKBOOK = HERE / "saas_company_reference_model.xlsx"
OUT = HERE / "expected.json"

SCALAR_COL = 9  # I
FIRST_PERIOD_COL = 10  # J
DATA_ROW = 3


def cell_value(v, typ):
    if v is None:
        raise ValueError("empty cell; recalculate the workbook before snapshotting")
    if typ.endswith("<bool>") or typ == "bool":
        return bool(v)
    return float(v)


def parse_tolerance(text):
    """'rel 1e-6' -> {"rel": 1e-6}; 'exact' -> {"exact": True}."""
    if text == "exact":
        return {"exact": True}
    kind, amount = text.split()
    if kind != "rel":
        raise ValueError(f"unknown tolerance {text!r}")
    return {"rel": float(amount)}


def read_rows(ws, n_values, first_col, meta_cols):
    """Rows keyed by id: metadata from the named columns plus a value or a list of values."""
    rows = {}
    for r in range(DATA_ROW, ws.max_row + 1):
        ident = ws.cell(r, 1).value
        if not ident:
            continue
        meta = {name: ws.cell(r, col).value for name, col in meta_cols.items()}
        typ = meta["type"]
        if typ.startswith(("series<", "annual<")):
            value = [cell_value(ws.cell(r, first_col + i).value, typ) for i in range(n_values)]
        else:
            value = cell_value(ws.cell(r, SCALAR_COL).value, typ)
        rows[ident] = {**meta, "value": value}
    return rows


def main():
    wb = load_workbook(WORKBOOK, data_only=True)

    tl = {wb["Timeline"].cell(r, 1).value: wb["Timeline"].cell(r, 2).value for r in range(2, 7)}
    timeline = {
        "start_date": tl["start_date"].date().isoformat() if isinstance(tl["start_date"], dt.datetime) else str(tl["start_date"]),
        "frequency": tl["frequency"],
        "periods": int(tl["periods"]),
        "months_per_period": int(tl["months_per_period"]),
        "fiscal_year_end": tl["fiscal_year_end"],
    }
    n = timeline["periods"]

    calcs_ws, annual_ws = wb["Calculations"], wb["Annual"]
    labels = [calcs_ws.cell(2, FIRST_PERIOD_COL + i).value for i in range(n)]
    years = [annual_ws.cell(1, SCALAR_COL + i).value for i in range(annual_ws.max_column - SCALAR_COL + 1)]
    years = [int(y) for y in years if y is not None]

    calc_meta = {"label": 2, "group": 3, "type": 4, "unit": 5, "formula": 6}
    rows = read_rows(calcs_ws, n, FIRST_PERIOD_COL, calc_meta)
    annual_meta = {"label": 2, "type": 3, "unit": 4, "formula": 5}
    annual = read_rows(annual_ws, len(years), SCALAR_COL, annual_meta)
    for ident, row in annual.items():
        row["group"] = "Annual"
        rows[ident] = row

    # Expected rows are links to the rows above; confirm the cached values agree.
    expected = []
    for sheet, n_values, first_col in (("Expected", n, FIRST_PERIOD_COL), ("Expected_Annual", len(years), SCALAR_COL)):
        ws = wb[sheet]
        meta = {"type": 3, "unit": 4, "tolerance": 6, "role": 7}
        for ident, row in read_rows(ws, n_values, first_col, meta).items():
            if row["value"] != rows[ident]["value"]:
                raise ValueError(f"{sheet} {ident} differs from its source row")
            expected.append({"id": ident, "role": row["role"], "tolerance": parse_tolerance(row["tolerance"])})

    rep = wb["Report"]
    report = [
        {"order": rep.cell(r, 1).value, "id": rep.cell(r, 2).value, "label": rep.cell(r, 3).value, "format": rep.cell(r, 4).value}
        for r in range(DATA_ROW, rep.max_row + 1)
        if rep.cell(r, 2).value
    ]

    out = {
        "model": "saas_company",
        "source": {
            "workbook": WORKBOOK.name,
            "sha256": hashlib.sha256(WORKBOOK.read_bytes()).hexdigest(),
            "tool": Path(__file__).name,
        },
        "timeline": timeline,
        "period_labels": labels,
        "years": years,
        "expected": expected,
        "report": report,
        "rows": rows,
    }
    OUT.write_text(json.dumps(out, indent=1) + "\n")
    print(f"wrote {OUT.name}: {len(rows)} rows, {len(expected)} expected, {len(report)} report lines")


if __name__ == "__main__":
    main()
