"""Build the reference model workbook from spec.py with live formulas."""
import datetime as dt
import re
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

import dsl
from spec import DRIVERS, SERIES, CALCS, ANNUAL, EXPECTED, REPORT, TIMELINE, PERIODS

OUT = sys.argv[1] if len(sys.argv) > 1 else "saas_company_reference_model.xlsx"

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_HEAD = Font(name=FONT, size=10, bold=True)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
F_SUB = Font(name=FONT, size=9, italic=True, color="666666")
F_MONO = Font(name="Consolas", size=9)
FILL_HEAD = PatternFill("solid", fgColor="E7E6E6")
BOTTOM = Border(bottom=Side(style="thin", color="999999"))

FIRST_COL = 10                      # J
SCALAR_COL = 9                      # I
COLS = [get_column_letter(FIRST_COL + i) for i in range(PERIODS)]
LAST = COLS[-1]
DATA_ROW = 3                        # rows 1-2 are headers on period sheets

FMT = {
    "currency": '$#,##0;($#,##0);"-"',
    "percent": '0.0%;(0.0%);"-"',
    "quantity": '#,##0.0;(#,##0.0);"-"',
    "count": '#,##0;(#,##0);"-"',
    "int": '#,##0;(#,##0);"-"',
    "float": '#,##0.0000;(#,##0.0000);"-"',
    "bool": "General",
    "date": "yyyy-mm-dd",
    "str": "General",
}


def base_type(typ):
    m = re.match(r"series<(\w+)>", typ)
    return m.group(1) if m else typ


def fmt_for(typ, unit=""):
    b = base_type(typ)
    if b == "float" and unit in ("days",):
        return '#,##0.00'
    return FMT[b]


def check_name(n):
    if re.fullmatch(r"[A-Za-z]{1,3}\d{1,7}", n) or n.lower() in ("r", "c"):
        raise ValueError(f"defined name {n} collides with a cell reference")


def header(ws, row, values, widths=None):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BOTTOM
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def period_headers(ws):
    """Row 1: period index t. Row 2: period label linked from Timeline."""
    c = ws.cell(row=1, column=SCALAR_COL, value="scalar")
    c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BOTTOM
    for i, col in enumerate(COLS):
        c = ws[f"{col}1"]
        c.value = i + 1
        c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BOTTOM
        c.alignment = Alignment(horizontal="right")
        c2 = ws[f"{col}2"]
        c2.value = f"=Timeline!{col}$13"
        c2.font = F_SUB
        c2.alignment = Alignment(horizontal="right")
        ws.column_dimensions[col].width = 13
    ws.cell(row=2, column=1, value="period label →").font = F_SUB


wb = Workbook()

# ------------------------------------------------------------------ README
rd = wb.active
rd.title = "README"
readme = [
    ("SaaS company reference model", F_HEAD),
    ("Reference model for the block model builder (acceptance test). Company-level, quarterly, 32 periods from 2027-01-01.", F_BASE),
    ("All figures are illustrative. It describes a generic subscription software company, not a real one.", F_BASE),
    ("", F_BASE),
    ("Sheets", F_HEAD),
    ("Timeline: timeline settings and per-period built-ins (t, period_start, period_end, year).", F_BASE),
    ("Drivers: scalar assumptions; each value cell is a defined name equal to its id. 'series' means see Driver_Series.", F_BASE),
    ("Driver_Series: time-varying drivers, one row per id.", F_BASE),
    ("Calculations: one row per derived line in dependency order. Column F is the DSL formula; the cells are the same logic as live Excel formulas.", F_BASE),
    ("Annual: fiscal-year rollups via annual().", F_BASE),
    ("Expected / Expected_Annual: acceptance outputs, linked live to Calculations and Annual. Compare at relative tolerance 1e-6.", F_BASE),
    ("Report: lines for the exported summary worksheet, in order, with display format (values linked for preview).", F_BASE),
    ("", F_BASE),
    ("Layout on period sheets", F_HEAD),
    ("Row 1 holds t (1 to 32); row 2 holds the period label; data starts in row 3. Column I holds scalar values; columns J to AO are periods 1 to 32.", F_BASE),
    ("", F_BASE),
    ("Color conventions", F_HEAD),
    ("Blue text: hardcoded inputs. Black text: formulas. Green text: links to another sheet.", F_BASE),
    ("", F_BASE),
    ("Conventions", F_HEAD),
    ("Percentages stored as fractions. Costs positive and subtracted explicitly. Flows per period; balances end of period. No circular references.", F_BASE),
    ("NPV convention: period t discounted by (1 + r)^t with t from 1, valued at the start of period 1 (Excel NPV).", F_BASE),
    ("See the accompanying notes file for the narrative, sources, simplifications and new functions needed.", F_BASE),
]
for i, (txt, f) in enumerate(readme, start=1):
    rd.cell(row=i, column=1, value=txt).font = f
rd.column_dimensions["A"].width = 140

# ------------------------------------------------------------------ Timeline
tl = wb.create_sheet("Timeline")
header(tl, 1, ["id", "value", "notes"], [22, 14, 60])
tl_rows = [
    ("start_date", dt.date.fromisoformat(TIMELINE["start_date"]), "First day of period 1.", "date"),
    ("frequency", TIMELINE["frequency"], "Period length.", "str"),
    ("periods", TIMELINE["periods"], "Number of periods.", "int"),
    ("months_per_period", TIMELINE["months_per_period"], "Follows from frequency (quarter = 3).", "int"),
    ("fiscal_year_end", TIMELINE["fiscal_year_end"], "Fiscal year equals calendar year.", "str"),
]
for i, (k, v, note, typ) in enumerate(tl_rows, start=2):
    tl.cell(row=i, column=1, value=k).font = F_BASE
    c = tl.cell(row=i, column=2, value=v)
    c.font, c.number_format = F_INPUT, FMT[typ] if typ != "str" else "General"
    tl.cell(row=i, column=3, value=note).font = F_BASE
    check_name(k)
    wb.defined_names[k] = DefinedName(k, attr_text=f"Timeline!$B${i}")

TL = {"t": 9, "period_start": 10, "period_end": 11, "year": 12, "label": 13}
tl.cell(row=8, column=1, value="Per-period built-ins (columns J onward are periods)").font = F_HEAD
for name, r in TL.items():
    tl.cell(row=r, column=1, value=name).font = F_BASE
for i, col in enumerate(COLS):
    prev = COLS[i - 1] if i else None
    tl[f"{col}{TL['t']}"] = 1 if i == 0 else f"={prev}{TL['t']}+1"
    tl[f"{col}{TL['period_start']}"] = "=start_date" if i == 0 else f"=EDATE({prev}{TL['period_start']},months_per_period)"
    tl[f"{col}{TL['period_end']}"] = f"=EDATE({col}{TL['period_start']},months_per_period)-1"
    tl[f"{col}{TL['year']}"] = f"=YEAR({col}{TL['period_end']})"
    tl[f"{col}{TL['label']}"] = f'=YEAR({col}{TL["period_start"]})&"Q"&ROUNDUP(MONTH({col}{TL["period_start"]})/3,0)'
    for r in TL.values():
        tl[f"{col}{r}"].font = F_BASE
    tl[f"{col}{TL['t']}"].font = F_INPUT if i == 0 else F_BASE
    tl[f"{col}{TL['period_start']}"].number_format = "yyyy-mm-dd"
    tl[f"{col}{TL['period_end']}"].number_format = "yyyy-mm-dd"
    tl[f"{col}{TL['year']}"].number_format = "0"
    tl.column_dimensions[col].width = 12
tl.freeze_panes = "B2"

# ------------------------------------------------------------------ Drivers
dr = wb.create_sheet("Drivers")
header(dr, 1, ["id", "label", "group", "type", "unit", "value", "low", "high", "source", "notes"],
       [34, 44, 18, 18, 20, 16, 14, 14, 60, 90])
refs = {}
for i, (ident, label, group, typ, unit, value, low, high, source, notes) in enumerate(DRIVERS, start=2):
    vals = [ident, label, group, typ, unit, value, low, high, source, notes]
    for j, v in enumerate(vals, start=1):
        c = dr.cell(row=i, column=j, value=v)
        c.font = F_BASE
    for j in (6, 7, 8):
        c = dr.cell(row=i, column=j)
        if c.value is not None and value != "series":
            c.font = F_INPUT
            c.number_format = fmt_for(typ, unit)
    if value == "series":
        dr.cell(row=i, column=6).font = F_SUB
    else:
        check_name(ident)
        wb.defined_names[ident] = DefinedName(ident, attr_text=f"Drivers!$F${i}")
        refs[ident] = ("scalar", ident)
dr.freeze_panes = "B2"
dr.auto_filter.ref = f"A1:J{len(DRIVERS) + 1}"

# ------------------------------------------------------------------ Driver_Series
ds = wb.create_sheet("Driver_Series")
header(ds, 1, ["id", "label", "group", "type", "unit", "source", "notes", ""], [30, 40, 16, 18, 12, 14, 40, 2])
ds.column_dimensions["I"].width = 8
period_headers(ds)
series_drivers = [d for d in DRIVERS if d[5] == "series"]
for k, d in enumerate(series_drivers):
    r = DATA_ROW + k
    ident, label, group, typ, unit, _, _, _, source, notes = d
    for j, v in enumerate([ident, label, group, typ, unit, source, "see Drivers"], start=1):
        ds.cell(row=r, column=j, value=v).font = F_BASE
    for i, col in enumerate(COLS):
        c = ds[f"{col}{r}"]
        c.value = SERIES[ident][i]
        c.font, c.number_format = F_INPUT, fmt_for(typ, unit)
    refs[ident] = ("series", "Driver_Series", r)
ds.freeze_panes = "J3"

# ------------------------------------------------------------------ Calculations
cs = wb.create_sheet("Calculations")
header(cs, 1, ["id", "label", "group", "type", "unit", "formula", "excel_ref", "notes"], [30, 40, 16, 18, 14, 60, 22, 50])
cs.column_dimensions["I"].width = 16
period_headers(cs)
calc_rows = {}
for k, row in enumerate(CALCS):
    r = DATA_ROW + k
    ident = row[0]
    calc_rows[ident] = r
    if row[3].startswith("series<"):
        refs[ident] = ("series", "Calculations", r)
    else:
        check_name(ident)
        wb.defined_names[ident] = DefinedName(ident, attr_text=f"Calculations!$I${r}")
        refs[ident] = ("scalar", ident)

ctx = dsl.ExcelCtx(refs, COLS, TL, "Calculations")
for k, (ident, label, group, typ, unit, formula, notes) in enumerate(CALCS):
    r = DATA_ROW + k
    series = typ.startswith("series<")
    excel_ref = f"Calculations!J{r}:{LAST}{r}" if series else f"Calculations!I{r}"
    for j, v in enumerate([ident, label, group, typ, unit, formula, excel_ref, notes], start=1):
        c = cs.cell(row=r, column=j, value=v)
        c.font = F_MONO if j == 6 else F_BASE
    node = dsl.parse(formula)
    fmt = fmt_for(typ, unit)
    if series:
        for col in COLS:
            c = cs[f"{col}{r}"]
            c.value = "=" + dsl.to_excel(node, ctx, col)
            c.font, c.number_format = F_BASE, fmt
    else:
        c = cs.cell(row=r, column=SCALAR_COL, value="=" + dsl.to_excel(node, ctx, None))
        c.font, c.number_format = F_BASE, fmt
cs.freeze_panes = "J3"
cs.auto_filter.ref = f"A1:H{DATA_ROW + len(CALCS) - 1}"

# ------------------------------------------------------------------ Annual
an = wb.create_sheet("Annual")
header(an, 1, ["id", "label", "type", "unit", "formula", "excel_ref", "notes", ""], [26, 20, 12, 10, 28, 22, 40, 2])
YEAR_COLS = [get_column_letter(SCALAR_COL + i) for i in range(PERIODS // 4)]
year_range = f"Timeline!$J${TL['year']}:${LAST}${TL['year']}"
for i, ycol in enumerate(YEAR_COLS):
    c = an[f"{ycol}1"]
    c.value = f"=MIN({year_range})" if i == 0 else f"={YEAR_COLS[i - 1]}1+1"
    c.font, c.fill, c.border, c.number_format = F_HEAD, FILL_HEAD, BOTTOM, "0"
    an.column_dimensions[ycol].width = 15
an.cell(row=2, column=1, value="fiscal year in row 1 →").font = F_SUB
annual_rows = {}
for k, (ident, label, typ, unit, formula) in enumerate(ANNUAL):
    r = DATA_ROW + k
    annual_rows[ident] = r
    node = dsl.parse(formula)
    src = node[2][0][1]
    sr = calc_rows[src]
    excel_ref = f"Annual!{YEAR_COLS[0]}{r}:{YEAR_COLS[-1]}{r}"
    for j, v in enumerate([ident, label, f"annual<{typ}>", unit, formula, excel_ref,
                           "Sum of the fiscal year's periods; output is on a yearly timeline."], start=1):
        an.cell(row=r, column=j, value=v).font = F_MONO if j == 5 else F_BASE
    for ycol in YEAR_COLS:
        c = an[f"{ycol}{r}"]
        c.value = f"=SUMIF({year_range},{ycol}$1,Calculations!$J{sr}:${LAST}{sr})"
        c.font, c.number_format = F_BASE, FMT[typ]
an.freeze_panes = "I3"

# ------------------------------------------------------------------ Expected
calc_by_id = {c[0]: c for c in CALCS}
ex = wb.create_sheet("Expected")
header(ex, 1, ["id", "label", "type", "unit", "source", "tolerance", "role", ""], [28, 40, 18, 14, 24, 12, 14, 2])
ex.column_dimensions["I"].width = 16
period_headers(ex)
REQUIRED = set(EXPECTED[:11])
for k, ident in enumerate(EXPECTED):
    r = DATA_ROW + k
    _, label, group, typ, unit, formula, notes = calc_by_id[ident]
    sr = calc_rows[ident]
    series = typ.startswith("series<")
    src_ref = f"Calculations!J{sr}:{LAST}{sr}" if series else f"Calculations!I{sr}"
    tol = "exact" if base_type(typ) == "bool" else "rel 1e-6"
    for j, v in enumerate([ident, label, typ, unit, src_ref, tol, "required" if ident in REQUIRED else "checkpoint"], start=1):
        ex.cell(row=r, column=j, value=v).font = F_BASE
    fmt = fmt_for(typ, unit)
    if series:
        for col in COLS:
            c = ex[f"{col}{r}"]
            c.value = f"=Calculations!{col}{sr}"
            c.font, c.number_format = F_LINK, fmt
    else:
        c = ex.cell(row=r, column=SCALAR_COL, value=f"=Calculations!I{sr}")
        c.font, c.number_format = F_LINK, fmt
ex.freeze_panes = "J3"

ea = wb.create_sheet("Expected_Annual")
header(ea, 1, ["id", "label", "type", "unit", "source", "tolerance", "role", ""], [26, 20, 18, 10, 24, 12, 14, 2])
for i, ycol in enumerate(YEAR_COLS):
    c = ea[f"{ycol}1"]
    c.value = f"=Annual!{ycol}1"
    c.font, c.fill, c.border, c.number_format = F_HEAD, FILL_HEAD, BOTTOM, "0"
    ea.column_dimensions[ycol].width = 15
for k, (ident, label, typ, unit, formula) in enumerate(ANNUAL):
    r = DATA_ROW + k
    ar = annual_rows[ident]
    for j, v in enumerate([ident, label, f"annual<{typ}>", unit, f"Annual!{YEAR_COLS[0]}{ar}:{YEAR_COLS[-1]}{ar}", "rel 1e-6", "checkpoint"], start=1):
        ea.cell(row=r, column=j, value=v).font = F_BASE
    for ycol in YEAR_COLS:
        c = ea[f"{ycol}{r}"]
        c.value = f"=Annual!{ycol}{ar}"
        c.font, c.number_format = F_LINK, FMT[typ]
ea.freeze_panes = "I3"

# ------------------------------------------------------------------ Report
rp = wb.create_sheet("Report")
header(rp, 1, ["order", "id", "label", "format", "source", "", "", ""], [7, 30, 40, 22, 24, 2, 2, 2])
rp.column_dimensions["I"].width = 16
period_headers(rp)
driver_series_ids = {d[0] for d in series_drivers}
labels = {c[0]: c[1] for c in CALCS} | {d[0]: d[1] for d in DRIVERS}
for k, (ident, fmt) in enumerate(REPORT):
    r = DATA_ROW + k
    if ident in driver_series_ids:
        sheet, sr, series = "Driver_Series", refs[ident][2], True
    else:
        sheet, sr, series = "Calculations", calc_rows[ident], calc_by_id[ident][3].startswith("series<")
    src_ref = f"{sheet}!J{sr}:{LAST}{sr}" if series else f"{sheet}!I{sr}"
    for j, v in enumerate([k + 1, ident, labels[ident], fmt, src_ref], start=1):
        rp.cell(row=r, column=j, value=v).font = F_BASE
    xfmt = fmt.replace("$#,##0", '$#,##0;($#,##0);"-"') if fmt == "$#,##0" else fmt
    if series:
        for col in COLS:
            c = rp[f"{col}{r}"]
            c.value = f"={sheet}!{col}{sr}"
            c.font, c.number_format = F_LINK, xfmt
    else:
        c = rp.cell(row=r, column=SCALAR_COL, value=f"={sheet}!I{sr}")
        c.font, c.number_format = F_LINK, xfmt
rp.freeze_panes = "J3"

for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 90

wb.save(OUT)
print("saved", OUT, "calcs", len(CALCS), "drivers", len(DRIVERS))
