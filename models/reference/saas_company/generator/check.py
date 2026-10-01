"""Compare recalculated workbook values with the numpy evaluation, row by row."""
import sys
import numpy as np
from openpyxl import load_workbook
import model_np
from spec import CALCS, ANNUAL, EXPECTED, PERIODS

path = sys.argv[1]
env, annual, years = model_np.run()
wb = load_workbook(path, data_only=True)
cs, an, ex, ea = wb["Calculations"], wb["Annual"], wb["Expected"], wb["Expected_Annual"]

REL, ABS = 1e-9, 1e-6
worst = (0.0, None)
fails = []


def cmp(ident, got, want, where):
    global worst
    if isinstance(want, (bool, np.bool_)):
        if bool(got) != bool(want):
            fails.append((where, ident, got, want))
        return
    if got is None:
        fails.append((where, ident, got, want)); return
    err = abs(float(got) - float(want))
    tol = max(ABS, REL * abs(float(want)))
    rel = err / max(abs(float(want)), 1.0)
    if rel > worst[0]:
        worst = (rel, f"{where} {ident}")
    if err > tol:
        fails.append((where, ident, got, want))


for k, row in enumerate(CALCS):
    r, ident, typ = 3 + k, row[0], row[3]
    assert cs.cell(r, 1).value == ident, (r, ident)
    if typ.startswith("series<"):
        for i in range(PERIODS):
            cmp(ident, cs.cell(r, 10 + i).value, env[ident][i], f"Calculations!t{i + 1}")
    else:
        cmp(ident, cs.cell(r, 9).value, env[ident], "Calculations!I")

for k, row in enumerate(ANNUAL):
    r, ident = 3 + k, row[0]
    assert an.cell(r, 1).value == ident
    for i, y in enumerate(years):
        assert an.cell(1, 9 + i).value == y
        cmp(ident, an.cell(r, 9 + i).value, annual[ident][i], f"Annual!{y}")
        cmp(ident, ea.cell(r, 9 + i).value, annual[ident][i], f"Expected_Annual!{y}")

calc_types = {c[0]: c[3] for c in CALCS}
for k, ident in enumerate(EXPECTED):
    r = 3 + k
    assert ex.cell(r, 1).value == ident
    if calc_types[ident].startswith("series<"):
        for i in range(PERIODS):
            cmp(ident, ex.cell(r, 10 + i).value, env[ident][i], f"Expected!t{i + 1}")
    else:
        cmp(ident, ex.cell(r, 9).value, env[ident], "Expected!I")

print("failures:", len(fails))
for f in fails[:20]:
    print("  ", f)
print("worst relative difference:", worst)
print("npv_value vs npv_check:", env["npv_value"], env["npv_check"])
