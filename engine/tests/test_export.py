"""Summary and export: CSV and XLSX carry the evaluated report rows, in the right tables and formats."""

import csv
import io

import pytest
from blockmodel.evaluate import evaluate
from blockmodel.export import ExportError, build_summary, export
from blockmodel.importers.rows import import_rows
from blockmodel.synthetic import synthetic_model
from openpyxl import load_workbook
from oracle import SPEC_PATH


@pytest.fixture(scope="module")
def reference():
    return import_rows(SPEC_PATH).model


def test_summary_splits_rows_by_shape(reference):
    summary = build_summary(reference)
    assert len(summary.per_period) == 28
    assert [r.label for r in summary.single] == ["NPV", "IRR (annualized)"]
    assert summary.sample
    npv = next(r for r in summary.single if r.format == "$#,##0")
    assert npv.values[0] == pytest.approx(51_239_654.553002805, rel=1e-9)


def test_csv_matches_evaluation(reference):
    content, media = export(reference, "csv")
    assert media.startswith("text/csv")
    rows = list(csv.reader(io.StringIO(content.decode())))
    header = next(r for r in rows if r[:2] == ["Line", "Unit"])
    assert header[2] == "2027Q1" and header[-1] == "Note"
    result = evaluate(reference)
    line = next(line for line in reference.report if line.label.startswith("Revenue"))
    want = result.outputs[str(line.block)][line.port].value
    got = next(r for r in rows if r and r[0] == line.label)
    assert [float(x) for x in got[2 : 2 + len(want)]] == pytest.approx(want, rel=1e-12)


def test_xlsx_has_three_sheets_with_values_and_formats(reference):
    content, _ = export(reference, "xlsx")
    book = load_workbook(io.BytesIO(content))
    assert book.sheetnames == ["Summary", "Assumptions", "Blocks"]
    sheet = book["Summary"]
    labels = {sheet.cell(r, 1).value: r for r in range(1, sheet.max_row + 1)}
    result = evaluate(reference)
    for line in reference.report:
        row = labels[line.label]
        want = result.outputs[str(line.block)][line.port].value
        want = want if isinstance(want, list) else [want]
        got = [sheet.cell(row, 3 + i).value for i in range(len(want))]
        assert got == pytest.approx(want, rel=1e-12), line.label
        assert sheet.cell(row, 3).number_format == (line.format or "General")
    assumptions = book["Assumptions"]
    assert assumptions.max_row - 1 == sum(
        b.spec in ("source.constant@1", "source.series@1") for b in reference.blocks
    )
    assert book["Blocks"].cell(2, 1).value == str(reference.blocks[0].uuid)


def test_empty_summary_is_an_error(reference):
    with pytest.raises(ExportError, match="Add to summary"):
        export(reference.model_copy(update={"report": []}), "csv")


@pytest.mark.parametrize("blocks, periods, rows", [(11, 4, 1), (250, 120, 10), (2000, 600, 10)])
def test_synthetic_models_export(blocks, periods, rows):
    """Synthetic models summarize their last layer, so they export like real models."""
    model = synthetic_model(blocks=blocks, periods=periods)
    summary = build_summary(model)
    assert [r.label for r in summary.per_period] == [f"End block {n}" for n in range(1, rows + 1)]
    assert all(r.note == "" and len(r.values) == periods for r in summary.per_period)
    content, media_type = export(model, "xlsx")
    assert media_type.endswith("spreadsheetml.sheet")
    sheet = load_workbook(io.BytesIO(content), read_only=True)["Summary"]
    assert any(row[0] == "End block 1" for row in sheet.iter_rows(values_only=True))
